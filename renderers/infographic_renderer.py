"""
renderers/infographic_renderer.py — Builds infographic slides using python-pptx shapes.

Renders timelines, process flows, and comparison layouts programmatically
using native shapes (NOT images) for maximum scalability and theme compliance.
"""

from __future__ import annotations

import logging
from typing import Any

from pptx import presentation as pptx_presentation
from pptx.util import Inches, Pt, Emu
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor

from config import (
    SLIDE_WIDTH_INCHES,
    SLIDE_HEIGHT_INCHES,
    MARGIN_LEFT_FRAC,
    MARGIN_RIGHT_FRAC,
    CONTENT_TOP_FRAC,
    CONTENT_HEIGHT_FRAC,
    MAX_TIMELINE_EVENTS,
    MAX_PROCESS_STEPS,
    FONT_SIZE_BODY,
    FONT_SIZE_SMALL,
    FONT_SIZE_FOOTNOTE,
)

logger = logging.getLogger(__name__)


def render_timeline(
    events: list[dict[str, Any]],
    slide: Any,
    accent_color: RGBColor | None = None,
    font_name: str = "Calibri",
) -> None:
    """Render a horizontal timeline infographic on the given slide.

    Args:
        events: List of event dicts with "year", "title", "description" keys.
        slide: python-pptx slide object.
        accent_color: Accent color from the master theme.
        font_name: Font family to use.
    """
    if not events:
        logger.warning("No events provided for timeline")
        return

    events = events[:MAX_TIMELINE_EVENTS]
    accent = accent_color or RGBColor(0x2E, 0x86, 0xAB)
    dark_text = RGBColor(0x33, 0x33, 0x33)
    light_text = RGBColor(0x66, 0x66, 0x66)

    # Layout dimensions
    sw = SLIDE_WIDTH_INCHES
    sh = SLIDE_HEIGHT_INCHES
    left_margin = sw * MARGIN_LEFT_FRAC
    right_margin = sw * MARGIN_RIGHT_FRAC
    usable_width = sw - left_margin - right_margin

    content_top = sh * CONTENT_TOP_FRAC
    content_height = sh * CONTENT_HEIGHT_FRAC

    # Timeline line position (center vertically in content area)
    line_y = content_top + content_height * 0.45
    line_thickness = Inches(0.04)

    # Draw the horizontal line
    line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(left_margin),
        Inches(line_y),
        Inches(usable_width),
        line_thickness,
    )
    line.fill.solid()
    line.fill.fore_color.rgb = accent
    line.line.fill.background()

    # Place event markers
    spacing = usable_width / max(len(events), 1)
    circle_size = 0.3

    for i, event in enumerate(events):
        cx = left_margin + spacing * (i + 0.5)
        above = (i % 2 == 0)  # Alternate above/below

        # Circle marker
        circle_left = cx - circle_size / 2
        circle_top = line_y - circle_size / 2 + 0.02
        circle = slide.shapes.add_shape(
            MSO_SHAPE.OVAL,
            Inches(circle_left),
            Inches(circle_top),
            Inches(circle_size),
            Inches(circle_size),
        )
        circle.fill.solid()
        circle.fill.fore_color.rgb = accent
        circle.line.fill.background()

        # Year label inside circle
        tf = circle.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = str(event.get("year", ""))
        run.font.size = Pt(8)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        run.font.name = font_name
        tf.paragraphs[0].space_before = Pt(0)
        tf.paragraphs[0].space_after = Pt(0)

        # Title and description text box
        text_width = spacing * 0.85
        text_height = 1.2

        if above:
            text_top = line_y - circle_size / 2 - text_height - 0.15
        else:
            text_top = line_y + circle_size / 2 + 0.25

        text_left = cx - text_width / 2

        # Clamp to slide bounds
        text_left = max(left_margin, min(text_left, sw - right_margin - text_width))

        txbox = slide.shapes.add_textbox(
            Inches(text_left),
            Inches(text_top),
            Inches(text_width),
            Inches(text_height),
        )
        tf = txbox.text_frame
        tf.word_wrap = True

        # Title
        p_title = tf.paragraphs[0]
        p_title.alignment = PP_ALIGN.CENTER
        run_title = p_title.add_run()
        run_title.text = event.get("title", "")
        run_title.font.size = Pt(11)
        run_title.font.bold = True
        run_title.font.color.rgb = dark_text
        run_title.font.name = font_name

        # Description
        desc = event.get("description", "")
        if desc:
            p_desc = tf.add_paragraph()
            p_desc.alignment = PP_ALIGN.CENTER
            run_desc = p_desc.add_run()
            run_desc.text = desc[:80]
            run_desc.font.size = Pt(9)
            run_desc.font.color.rgb = light_text
            run_desc.font.name = font_name

    logger.info("Rendered timeline with %d events", len(events))


def render_process_flow(
    steps: list[dict[str, Any]],
    slide: Any,
    direction: str = "horizontal",
    accent_color: RGBColor | None = None,
    font_name: str = "Calibri",
) -> None:
    """Render a process flow infographic on the given slide.

    Args:
        steps: List of step dicts with "number", "title", "description".
        slide: python-pptx slide object.
        direction: "horizontal" or "vertical".
        accent_color: Accent color from the master theme.
        font_name: Font family to use.
    """
    if not steps:
        logger.warning("No steps provided for process flow")
        return

    steps = steps[:MAX_PROCESS_STEPS]
    accent = accent_color or RGBColor(0x2E, 0x86, 0xAB)
    dark_text = RGBColor(0x33, 0x33, 0x33)
    white = RGBColor(0xFF, 0xFF, 0xFF)
    light_bg = RGBColor(0xF5, 0xF5, 0xF5)

    sw = SLIDE_WIDTH_INCHES
    sh = SLIDE_HEIGHT_INCHES
    left_margin = sw * MARGIN_LEFT_FRAC
    right_margin = sw * MARGIN_RIGHT_FRAC
    usable_width = sw - left_margin - right_margin

    content_top = sh * CONTENT_TOP_FRAC
    content_height = sh * CONTENT_HEIGHT_FRAC

    n = len(steps)

    if direction == "horizontal" and n <= MAX_PROCESS_STEPS:
        # Determine layout: single row or two rows
        if n <= 4:
            rows = [steps]
        else:
            mid = (n + 1) // 2
            rows = [steps[:mid], steps[mid:]]

        row_height = content_height / len(rows)

        for row_idx, row_steps in enumerate(rows):
            cols = len(row_steps)
            box_width = min(2.5, (usable_width - (cols - 1) * 0.5) / cols)
            box_height = row_height * 0.7
            gap = (usable_width - cols * box_width) / max(cols - 1, 1) if cols > 1 else 0
            gap = min(gap, 1.0)
            total_w = cols * box_width + (cols - 1) * gap
            start_x = left_margin + (usable_width - total_w) / 2
            row_top = content_top + row_idx * row_height + (row_height - box_height) / 2

            for i, step in enumerate(row_steps):
                bx = start_x + i * (box_width + gap)

                # Rounded rectangle
                shape = slide.shapes.add_shape(
                    MSO_SHAPE.ROUNDED_RECTANGLE,
                    Inches(bx),
                    Inches(row_top),
                    Inches(box_width),
                    Inches(box_height),
                )
                shape.fill.solid()
                shape.fill.fore_color.rgb = light_bg
                shape.line.color.rgb = accent
                shape.line.width = Pt(1.5)

                # Step number circle
                circle_size = 0.42
                circle = slide.shapes.add_shape(
                    MSO_SHAPE.OVAL,
                    Inches(bx + box_width / 2 - circle_size / 2),
                    Inches(row_top - circle_size / 2),
                    Inches(circle_size),
                    Inches(circle_size),
                )
                circle.fill.solid()
                circle.fill.fore_color.rgb = accent
                circle.line.fill.background()

                tf_circle = circle.text_frame
                tf_circle.word_wrap = False
                p = tf_circle.paragraphs[0]
                p.alignment = PP_ALIGN.CENTER
                run = p.add_run()
                run.text = str(step.get("number", i + 1))
                run.font.size = Pt(14)
                run.font.bold = True
                run.font.color.rgb = white
                run.font.name = font_name

                # Title and description inside the box
                tf = shape.text_frame
                tf.word_wrap = True
                tf.paragraphs[0].space_before = Pt(16)

                p_title = tf.paragraphs[0]
                p_title.alignment = PP_ALIGN.CENTER
                run_title = p_title.add_run()
                run_title.text = step.get("title", "")
                run_title.font.size = Pt(12)
                run_title.font.bold = True
                run_title.font.color.rgb = dark_text
                run_title.font.name = font_name

                desc = step.get("description", "")
                if desc:
                    p_desc = tf.add_paragraph()
                    p_desc.alignment = PP_ALIGN.CENTER
                    p_desc.space_before = Pt(4)
                    run_desc = p_desc.add_run()
                    run_desc.text = desc[:60]
                    run_desc.font.size = Pt(9)
                    run_desc.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
                    run_desc.font.name = font_name

                # Arrow to next step (if not the last in the row)
                if i < cols - 1:
                    arrow_left = bx + box_width
                    arrow_top = row_top + box_height / 2 - 0.08
                    arrow_width = gap
                    arrow = slide.shapes.add_shape(
                        MSO_SHAPE.RIGHT_ARROW,
                        Inches(arrow_left),
                        Inches(arrow_top),
                        Inches(arrow_width * 0.6),
                        Inches(0.2),
                    )
                    arrow.fill.solid()
                    arrow.fill.fore_color.rgb = accent
                    arrow.line.fill.background()
    else:
        # Vertical layout
        box_width = usable_width * 0.5
        box_height = min(1.0, (content_height - (n - 1) * 0.2) / n)
        start_x = left_margin + (usable_width - box_width) / 2

        for i, step in enumerate(steps):
            by = content_top + i * (box_height + 0.3)

            shape = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE,
                Inches(start_x),
                Inches(by),
                Inches(box_width),
                Inches(box_height),
            )
            shape.fill.solid()
            shape.fill.fore_color.rgb = light_bg
            shape.line.color.rgb = accent
            shape.line.width = Pt(1.5)

            # Number circle
            circle_size = 0.4
            circle = slide.shapes.add_shape(
                MSO_SHAPE.OVAL,
                Inches(start_x - circle_size / 2),
                Inches(by + box_height / 2 - circle_size / 2),
                Inches(circle_size),
                Inches(circle_size),
            )
            circle.fill.solid()
            circle.fill.fore_color.rgb = accent
            circle.line.fill.background()

            tf_c = circle.text_frame
            p = tf_c.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            run = p.add_run()
            run.text = str(step.get("number", i + 1))
            run.font.size = Pt(12)
            run.font.bold = True
            run.font.color.rgb = white
            run.font.name = font_name

            # Text in the box
            tf = shape.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT
            run = p.add_run()
            run.text = step.get("title", "")
            run.font.size = Pt(13)
            run.font.bold = True
            run.font.color.rgb = dark_text
            run.font.name = font_name

            desc = step.get("description", "")
            if desc:
                p2 = tf.add_paragraph()
                p2.alignment = PP_ALIGN.LEFT
                run2 = p2.add_run()
                run2.text = desc[:80]
                run2.font.size = Pt(10)
                run2.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
                run2.font.name = font_name

            # Down arrow
            if i < n - 1:
                arrow = slide.shapes.add_shape(
                    MSO_SHAPE.DOWN_ARROW,
                    Inches(start_x + box_width / 2 - 0.12),
                    Inches(by + box_height),
                    Inches(0.24),
                    Inches(0.2),
                )
                arrow.fill.solid()
                arrow.fill.fore_color.rgb = accent
                arrow.line.fill.background()

    logger.info("Rendered process flow with %d steps (%s)", len(steps), direction)


def render_comparison(
    left_label: str,
    right_label: str,
    dimensions: list[dict[str, Any]],
    slide: Any,
    accent_color: RGBColor | None = None,
    font_name: str = "Calibri",
) -> None:
    """Render a comparison infographic with two columns and a center divider.

    Args:
        left_label: Label for the left column.
        right_label: Label for the right column.
        dimensions: List of dicts with "aspect", "left", "right" keys.
        slide: python-pptx slide object.
        accent_color: Accent color from the master theme.
        font_name: Font family to use.
    """
    if not dimensions:
        logger.warning("No dimensions provided for comparison")
        return

    accent = accent_color or RGBColor(0x2E, 0x86, 0xAB)
    dark_text = RGBColor(0x33, 0x33, 0x33)
    white = RGBColor(0xFF, 0xFF, 0xFF)
    light_bg = RGBColor(0xF5, 0xF7, 0xFA)
    alt_bg = RGBColor(0xE8, 0xEC, 0xF1)

    sw = SLIDE_WIDTH_INCHES
    sh = SLIDE_HEIGHT_INCHES
    left_margin = sw * MARGIN_LEFT_FRAC
    right_margin = sw * MARGIN_RIGHT_FRAC
    usable_width = sw - left_margin - right_margin

    content_top = sh * CONTENT_TOP_FRAC
    content_height = sh * CONTENT_HEIGHT_FRAC

    # Three columns: left | center (aspect labels) | right
    center_width = usable_width * 0.2
    side_width = (usable_width - center_width) / 2

    left_x = left_margin
    center_x = left_margin + side_width
    right_x = center_x + center_width

    # Column headers
    header_height = 0.6

    # Left header
    left_header = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(left_x),
        Inches(content_top),
        Inches(side_width),
        Inches(header_height),
    )
    left_header.fill.solid()
    left_header.fill.fore_color.rgb = accent
    left_header.line.fill.background()
    tf = left_header.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = left_label
    run.font.size = Pt(16)
    run.font.bold = True
    run.font.color.rgb = white
    run.font.name = font_name

    # Right header
    right_header = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(right_x),
        Inches(content_top),
        Inches(side_width),
        Inches(header_height),
    )
    right_header.fill.solid()
    right_header.fill.fore_color.rgb = accent
    right_header.line.fill.background()
    tf = right_header.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = right_label
    run.font.size = Pt(16)
    run.font.bold = True
    run.font.color.rgb = white
    run.font.name = font_name

    # Center header ("VS" or aspect label)
    center_header = slide.shapes.add_shape(
        MSO_SHAPE.OVAL,
        Inches(center_x + center_width / 2 - 0.3),
        Inches(content_top + header_height / 2 - 0.3),
        Inches(0.6),
        Inches(0.6),
    )
    center_header.fill.solid()
    center_header.fill.fore_color.rgb = dark_text
    center_header.line.fill.background()
    tf = center_header.text_frame
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = "VS"
    run.font.size = Pt(12)
    run.font.bold = True
    run.font.color.rgb = white
    run.font.name = font_name

    # Dimension rows
    row_top = content_top + header_height + 0.15
    max_dims = min(len(dimensions), 6)
    row_height = min(0.7, (content_height - header_height - 0.3) / max_dims)

    for i, dim in enumerate(dimensions[:max_dims]):
        ry = row_top + i * (row_height + 0.08)
        bg = light_bg if i % 2 == 0 else alt_bg

        # Left cell
        left_cell = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(left_x),
            Inches(ry),
            Inches(side_width),
            Inches(row_height),
        )
        left_cell.fill.solid()
        left_cell.fill.fore_color.rgb = bg
        left_cell.line.fill.background()
        tf = left_cell.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = dim.get("left", "")
        run.font.size = Pt(12)
        run.font.color.rgb = dark_text
        run.font.name = font_name

        # Center cell (aspect label)
        center_cell = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(center_x),
            Inches(ry),
            Inches(center_width),
            Inches(row_height),
        )
        center_cell.fill.solid()
        center_cell.fill.fore_color.rgb = RGBColor(0xE3, 0xE8, 0xEF)
        center_cell.line.fill.background()
        tf = center_cell.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = dim.get("aspect", "")
        run.font.size = Pt(11)
        run.font.bold = True
        run.font.color.rgb = accent
        run.font.name = font_name

        # Right cell
        right_cell = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(right_x),
            Inches(ry),
            Inches(side_width),
            Inches(row_height),
        )
        right_cell.fill.solid()
        right_cell.fill.fore_color.rgb = bg
        right_cell.line.fill.background()
        tf = right_cell.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = dim.get("right", "")
        run.font.size = Pt(12)
        run.font.color.rgb = dark_text
        run.font.name = font_name

    logger.info("Rendered comparison with %d dimensions", len(dimensions[:max_dims]))
